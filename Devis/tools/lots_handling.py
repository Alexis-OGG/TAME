from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ============================================================
# Sommaire Lots Création
# ============================================================

def dessiner_diagramme_lots(slide, placeholder, liste_lots):
    """
    Dessine le diagramme des phases et des lots dans l'espace réservé.
    :param liste_lots: Liste d'objets issus de votre classe Pydantic 'Lot'
    """
    
    # 1. Regroupement des lots par phase
    # On crée un dictionnaire : {"Phase 1": [lot1, lot2], "Phase 2": [lot3]}
    groupes_phases = {}
    for lot in liste_lots:
        if lot.phase not in groupes_phases:
            groupes_phases[lot.phase] = []
        groupes_phases[lot.phase].append(lot)
        
    # 2. Coordonnées de base (issues du placeholder)
    zone_x = placeholder.left
    zone_y = placeholder.top
    zone_largeur = placeholder.width
    
    # Nettoyage du placeholder vide
    element_xml = placeholder.element
    element_xml.getparent().remove(element_xml)
    
    # 3. Paramètres de dimensionnement et couleurs
    largeur_gauche = zone_largeur * 0.25      # Le bloc phase prend 25% de la largeur
    espace_milieu = zone_largeur * 0.03       # Espace de 3% entre gauche et droite
    largeur_droite = zone_largeur * 0.72      # Les lots prennent le reste (72%)
    
    hauteur_lot = Cm(1.2)       # Hauteur d'une barre de lot
    espace_lot = Cm(0.2)        # Espace vertical entre chaque lot
    espace_phase = Cm(0.5)      # Espace vertical entre chaque grand bloc de phase
    
    # Palette de couleurs extraite de votre image
    couleur_fond_phase = RGBColor(234, 230, 223) # Beige clair
    couleur_fond_lot = RGBColor(226, 235, 244)   # Bleu/Gris très clair
    couleur_texte = RGBColor(38, 74, 127)        # Bleu foncé/Gris
    
    pos_y_courante = zone_y
    
    # 4. Boucle de dessin
    for nom_phase, lots_de_la_phase in groupes_phases.items():
        
        nb_lots = len(lots_de_la_phase)
        
        # La hauteur du grand bloc de gauche dépend du nombre de lots à droite
        hauteur_phase_totale = (nb_lots * hauteur_lot) + ((nb_lots - 1) * espace_lot)
        
        # --- A. DESSIN DU GROS BLOC GAUCHE (PHASE) ---
        bloc_gauche = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            zone_x, pos_y_courante, largeur_gauche, hauteur_phase_totale
        )
        
        # Style du bloc gauche
        bloc_gauche.fill.solid()
        bloc_gauche.fill.fore_color.rgb = couleur_fond_phase
        bloc_gauche.line.fill.background() # Pas de bordure
        
        # Texte du bloc gauche
        tf_gauche = bloc_gauche.text_frame
        tf_gauche.word_wrap = True
        p_gauche = tf_gauche.paragraphs[0]
        p_gauche.text = nom_phase
        p_gauche.font.bold = True
        p_gauche.font.size = Pt(12)
        p_gauche.font.color.rgb = couleur_texte
        p_gauche.alignment = PP_ALIGN.CENTER
        
        # --- B. DESSIN DES BLOCS DROITS (LOTS) ---
        pos_y_lot = pos_y_courante
        
        for lot in lots_de_la_phase:
            
            x_droite = zone_x + largeur_gauche + espace_milieu
            
            # 1. Le fond du bloc (Rectangle arrondi)
            bloc_droit = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                x_droite, pos_y_lot, largeur_droite, hauteur_lot
            )
            bloc_droit.fill.solid()
            bloc_droit.fill.fore_color.rgb = couleur_fond_lot
            bloc_droit.line.fill.background()
            
            # 2. Le texte du lot (intégré directement dans la forme)
            tf_droit = bloc_droit.text_frame
            # MAGIE ICI : On ajoute une marge à gauche dans la boîte de texte
            # pour laisser la place à l'icône, sans que le texte ne se superpose !
            tf_droit.margin_left = Cm(1.5) 
            
            p_droit = tf_droit.paragraphs[0]
            p_droit.text = lot.nom.upper() # Mis en majuscule comme sur l'image
            p_droit.font.bold = True
            p_droit.font.size = Pt(11)
            p_droit.font.color.rgb = couleur_texte
            p_droit.alignment = PP_ALIGN.LEFT
            
            # 3. L'icône du lot
            if lot.image_lot and lot.image_lot.strip() != "":
                try:
                    # On insère l'image par-dessus le bloc droit
                    slide.shapes.add_picture(
                        lot.image_lot, 
                        x_droite + Cm(0.3),   # Marge interne de l'image
                        pos_y_lot + Cm(0.2),  # Centrage vertical approximatif
                        height=Cm(0.8)        # L'image est plus petite que la barre (1.2cm)
                    )
                except Exception as e:
                    print(f"⚠️ Image {lot.image_lot} introuvable : {e}")
            
            # On descend pour le prochain lot
            pos_y_lot += hauteur_lot + espace_lot
            
        # On descend pour la prochaine Phase
        pos_y_courante += hauteur_phase_totale + espace_phase


# ============================================================
# AJOUT DONNEES/LIVRABLES
# ============================================================

def dessiner_documents_lot(slide, placeholder, documents_lot):
    """
    Dessine les Données d'entrée et Livrables dans un placeholder
    à partir d'une instance du modèle Pydantic DocumentsLot.
    """
    
    # 1. Récupérer les coordonnées de la zone réservée
    zone_x = placeholder.left
    zone_y = placeholder.top
    zone_largeur = placeholder.width
    
    # 2. Supprimer le placeholder vide
    element_xml = placeholder.element
    element_xml.getparent().remove(element_xml)
    
    pos_y_courante = zone_y
    decalage_icone = Cm(0.2)
    decalage_texte = Cm(1.2)
    bleu_titre = RGBColor(79, 113, 172)
    
    # 3. Préparer une liste pour boucler proprement sur vos deux catégories
    sections = []
    
    # On n'ajoute la section que si elle contient des documents
    if documents_lot.Donnees_entree:
        sections.append(("Données d'entrée", documents_lot.Donnees_entree))
        
    if documents_lot.Livrables:
        sections.append(("Livrables", documents_lot.Livrables))
        
    # 4. Boucle de dessin
    for titre_section, liste_documents in sections:
        
        # --- DESSIN DU TITRE DE LA SECTION ---
        titre_box = slide.shapes.add_textbox(zone_x, pos_y_courante, zone_largeur, Cm(1))
        tf = titre_box.text_frame
        tf.word_wrap = True 
        
        p = tf.paragraphs[0]
        p.text = titre_section

        p.font.bold = True
        p.font.size = Pt(16) 
        p.font.color.rgb = bleu_titre
        
        pos_y_courante += Cm(1.0)
        
        # --- DESSIN DES DOCUMENTS ---
        for doc in liste_documents:
            
            # A. Icône (on vérifie si le champ n'est pas vide)
            if doc.img_doc and doc.img_doc.strip() != "":
                try:
                    slide.shapes.add_picture(
                        doc.img_doc, 
                        zone_x + decalage_icone, 
                        pos_y_courante, 
                        height=Cm(0.8)
                    )
                except Exception as e:
                    print(f"⚠️ Impossible d'insérer l'icône '{doc.img_doc}': {e}")
                    
            # B. Texte du document
            element_box = slide.shapes.add_textbox(
                zone_x + decalage_texte, 
                pos_y_courante + Cm(0.1), 
                zone_largeur - decalage_texte, 
                Cm(0.8)
            )
            tf_elem = element_box.text_frame
            tf_elem.word_wrap = True # Pour les noms de docs très longs
            
            p_elem = tf_elem.paragraphs[0]
            p_elem.text = doc.text_doc
            p_elem.font.size = Pt(11)
            
            # On descend pour le prochain document
            pos_y_courante += Cm(0.9)
            
        # Marge d'espacement avant de passer aux Livrables (si existants)
        pos_y_courante += Cm(0.6)

if __name__ == "__main__" :
    pass