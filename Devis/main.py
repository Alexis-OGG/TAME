
#Import modules
from classes import *
from tools.generate_table import dessiner_tableau_dynamique
from tools.generate_gantt import ajouter_gantt
from tools.lots_handling import dessiner_documents_lot, dessiner_diagramme_lots
from tools.feed_hypothesis import remplir_placeholder_hypotheses
from tools.text_handling import injecter_texte_auto_ajuste

#Import librairies
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import re
from dotenv import load_dotenv



# ============================================================
# CONSTRUCTION DU POWERPOINT
# ============================================================

def main(payload:PresentationData):

    #Source du Template
    prs = Presentation("input/DEV_Template.pptx")

    # Ajout Introduction
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("1_Introduction"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
        dico[f"{index}"] = placeholder.placeholder_format.idx

    slide.placeholders[dico["0"]].text = payload.titre_presentation
    #slide.placeholders[dico["1"]].insert_picture(payload.logo_client)
    slide.placeholders[dico["3"]].text = payload.nom_projet
    slide.placeholders[dico["2"]].text = "DEV-"+payload.id_projet+"-"+payload.Annee

    #AJout Slide Présentation TAME
    slide=prs.slides.add_slide(prs.slide_layouts.get_by_name("2_Introduction"))
    slide=prs.slides.add_slide(prs.slide_layouts.get_by_name("3_Introduction"))
    slide=prs.slides.add_slide(prs.slide_layouts.get_by_name("4_Introduction"))

    #Ajout Transition
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Transition"))
    for placeholder in slide.placeholders:
        placeholder.text = "Besoin"

    #Ajout Besoin
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Besoin"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
        dico[f"{index}"] = placeholder.placeholder_format.idx

    injecter_texte_auto_ajuste(slide.placeholders[dico["0"]],payload.Besoin.Avancement)
    injecter_texte_auto_ajuste(slide.placeholders[dico["1"]],payload.Besoin.Besoin)
    injecter_texte_auto_ajuste(slide.placeholders[dico["2"]],payload.Besoin.Role_TAME)
    # slide.placeholders[dico["0"]].text_frame.paragraphs[0].text = payload.Besoin.Avancement
    
    # slide.placeholders[dico["1"]].text_frame.paragraphs[0].text = payload.Besoin.Besoin
    # slide.placeholders[dico["2"]].text_frame.paragraphs[0].text = payload.Besoin.Role_TAME

    #Ajout Données d'entrée
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Entrees"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
        dico[f"{index}"] = placeholder.placeholder_format.idx
    slide.placeholders[dico["0"]].text = "Données d'entrée"
    slide.placeholders[dico["1"]].text= payload.donnees_entree

    #Ajout les hypothèses
    for hyp in payload.hypotheses:
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("1_Hypotheses"))
        dico={}
        for index,placeholder in enumerate(slide.placeholders):
            dico[f"{index}"] = placeholder.placeholder_format.idx
        slide.placeholders[dico["0"]].text = hyp.Type_Hypothese
        remplir_placeholder_hypotheses(slide.placeholders[dico["1"]], hyp.List_Hypotheses)

    #Ajout SOmmaire Lot
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Sommaire_lots"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
        dico[f"{index}"] = placeholder.placeholder_format.idx  
    dessiner_diagramme_lots(slide, slide.placeholders[dico["0"]], payload.lots_list.lots)
    slide.placeholders[dico["1"]].text = payload.entreprise
    #AJout de chaque diapo Lot
    for lot in payload.lots_list.lots:
            slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Lot"))
            dico={}
            for index,placeholder in enumerate(slide.placeholders):
                dico[f"{index}"] = placeholder.placeholder_format.idx
            remplir_placeholder_hypotheses(slide.placeholders[dico["3"]], lot.details_lot)
            slide.placeholders[dico["0"]].text = lot.nom
            slide.shapes.add_picture(lot.image_lot,slide.placeholders[dico["1"]].left,slide.placeholders[dico["1"]].top,slide.placeholders[dico["1"]].width*0.85,slide.placeholders[dico["1"]].height*0.85)
            slide.placeholders[dico["1"]].element.getparent().remove(slide.placeholders[dico["1"]].element)
            dessiner_documents_lot(slide,slide.placeholders[dico["2"]],lot.documents)

    #Ajout Transition
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Transition"))
    for placeholder in slide.placeholders:
        placeholder.text = "Synthèse"

    #Ajout Synthèse
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Synthese_finance"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
        dico[f"{index}"] = placeholder.placeholder_format.idx
    slide.placeholders[dico["1"]].text = "Synthèse"
    dessiner_tableau_dynamique(slide,slide.placeholders[dico["0"]],payload.lots_list)

    #Ajout du Planning
    if payload.id_semaine_debut :
        debut_gantt_init= payload.id_semaine_debut
        #TRi des lots par phase
        tri_lots_phases = {}
        for lot in payload.lots_list.lots:
            if not tri_lots_phases.get(lot.phase,{}):
                tri_lots_phases[lot.phase]={"lots":[lot],"nb_semaines":lot.duree+lot.debut,"debut":lot.debut}
            else:
                tri_lots_phases[lot.phase]["lots"].append(lot) 
                tri_lots_phases[lot.phase]["nb_semaines"] =lot.duree +lot.debut
        #Création d"une slide planning par phase
        
        for nom_phase,liste_lot in tri_lots_phases.items():

            ajouter_gantt(
                    annee = payload.Annee,
                    prs = prs,
                    layout=prs.slide_layouts.get_by_name("Planning"),
                    title = f"Planning - {nom_phase}",
                    lots=liste_lot["lots"],
                    nombre_semaines_tot=liste_lot["debut"],
                    nombre_semaines_phase=liste_lot["nb_semaines"]-liste_lot["debut"],
                    debut_gantt=debut_gantt_init,
            )


    #Ajout Transition
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Transition"))
    for placeholder in slide.placeholders:
        placeholder.text = "Annexes"

    #Ajout propriété intellectuelles
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Prop_int"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
            dico[f"{index}"] = placeholder.placeholder_format.idx  
    slide.placeholders[dico["1"]].text = "PROPRIETE INTELLECTUELLE"

    text =f"""Domaine de l'offre
\tLe DOMAINE couvert par cette offre est « [COMPLETER] ».
Connaissances propres
\tChaque Partie reste propriétaire de ses Connaissances Propres ou antérieures. À ce titre, chaque Partie reste libre d'exploiter ses Connaissances Propres ou antérieures comme bon lui semble. Lorsqu'elle emploie pour l'exécution du projet ses Connaissances Propres ou antérieures, TRONICO concède à {payload.entreprise}, sans frais additionnel au prix de la Commande, une licence d'exploitation des droits afférents. Cette licence sur les Connaissances Propres ou antérieures est concédée pour permettre à {payload.entreprise} de jouir pleinement des droits dont elle dispose sur les Résultats conformément aux dispositions de l'article « Résultats » de la présente proposition. Cette licence est concédée uniquement pour les Connaissances Propres ou antérieures qui font parties des Résultats.

Résultats
\tLes droits de propriété intellectuelle attachés aux Résultats ainsi que les livrables associés sont la propriété de {payload.entreprise} après paiement des factures dues à TRONICO. TRONICO bénéficie d'une licence non-exclusive d'exploitation des Résultats pour satisfaire tous besoins de son choix ou toute demande d'un autre client en dehors du Domaine.
"""
    text_frame = slide.placeholders[dico["0"]].text_frame
    for i, ligne in enumerate(text.split('\n')):
        if not ligne.strip():
            continue # Ignore les lignes vides
            
        niveau = ligne.count('\t')
        texte_propre = ligne.strip()
        
        if len(text_frame.paragraphs) == 0 or (i == 0 and text_frame.paragraphs[0].text == ""):
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        if niveau == 1 :
            p.font.name = "Poppins Light"
        p.level = niveau
        p.alignment = PP_ALIGN.JUSTIFY

        morceaux = re.split(r'(\[[^\]]+\])', texte_propre)
        
        for morceau in morceaux:
            if not morceau:
                continue
            run = p.add_run()
            run.text = morceau
            
            # Si le morceau est entouré de crochets, on le colore en rouge
            if morceau.startswith('[') and morceau.endswith(']'):
                run.font.color.rgb = RGBColor(255, 0, 0)
                run.font.bold = True

    text = f"""Non-sollicitation
\t{payload.entreprise} s'interdit expressément, pendant la durée du projet défini dans la présente proposition et vingt-quatre (24) mois après son expiration, de solliciter en vue d'une embauche ou d'embaucher directement ou indirectement tout membre du personnel de TRONICO.  
\tEn cas d'infraction à la présente interdiction, {payload.entreprise} sera tenue de payer immédiatement à TRONICO, à titre de clause pénale, une indemnité forfaitaire d'un montant égal à six (6) mois du dernier salaire brut mensuel de la personne sollicitée ou embauchée, majorée de tous les frais de recrutement d'un remplaçant.

Validité
\tCette proposition est valide pour une durée de 1 mois.
\tTRONICO ne peut être tenue pour responsable de toute anomalie consécutive à une information qu'elle n'aurait pas reçue. Les informations prises en compte sont la présente offre technique, le document de référence, et d'éventuels documents remis au cours de la réunion de lancement d'affaire.

Garantie
\tLa période de garantie pour les défauts de conception ou de fabrication de matériel est d'un an à compter de la date de livraison. 
\tLa garantie ne s'applique pas aux cartes d'évaluation et POC et si {payload.entreprise} :
\t\tutilise mal les produits et l'équipement, 
\t\test négligent dans le stockage, l'entretien ou l'utilisation des produits et de l'équipement, 
\t\tmodifie ou fait modifier par des tiers les produits ou les équipements, 
\t\teffectue ou fait effectuer par des tiers des réparations ou des mises à niveau des produits, à moins que le vendeur n'y ait expressément consenti.

Conditions de paiement
\t[COMPLETER]% du montant global du projet facturés à la commande.
\tPaiement à [COMPLETER]
\tTronico étant agrémenté CIR, les lots de développement définis dans cette proposition peuvent donner droit à des crédits d'impôt après facturation.



"""
    #Ajout condictions générales de ventes
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Cond_vente"))
    dico={}
    for index,placeholder in enumerate(slide.placeholders):
            dico[f"{index}"] = placeholder.placeholder_format.idx  
    slide.placeholders[dico["1"]].text = "CONDITIONS générales de l'étude"
    text_frame = slide.placeholders[dico["0"]].text_frame
    for i, ligne in enumerate(text.split('\n')):
            if not ligne.strip():
                continue # Ignore les lignes vides
                
            niveau = ligne.count('\t')
            texte_propre = ligne.strip()
            
            if len(text_frame.paragraphs) == 0 or (i == 0 and text_frame.paragraphs[0].text == ""):
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            if niveau == 1 :
                p.font.name = "Poppins Light"
            p.level = niveau
            p.alignment = PP_ALIGN.JUSTIFY

            morceaux = re.split(r'(\[[^\]]+\])', texte_propre)

            for morceau in morceaux:
                if not morceau:
                    continue
                run = p.add_run()
                run.text = morceau
                
                # Si le morceau est entouré de crochets, on le colore en rouge
                if morceau.startswith('[') and morceau.endswith(']'):
                    run.font.color.rgb = RGBColor(255, 0, 0)
                    run.font.bold = True
    

    #Ajout Conditions générales de vente
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("1_CGDV"))
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("2_CGDV"))

    #Sauvegarde
    prs.save('output/output.pptx')


#Zone de tests
if __name__ == "__main__" :
    load_dotenv()
    #Exemple ICONEUS
    # test = PresentationData(
    #         titre_presentation="PROPOSITION TECHNIQUE & FINANCIÈRE",
    #         logo_client = "img/Image2.png",
    #         nom_projet="ICONEUS - Multiplexer probe",
    #         Annee="23/07/2026",
    #         entreprise="ICONEUS",
    #         id_projet="000970_01",
    #         donnees_entree = "Cdc",
    #         Besoin =Besoin(
    #             Avancement= "ICONEUS à transmis un cahier des charges présenté aux équipes le 24 juin. Des échanges au travers d'un Tame-Carefichier questions / réponses ont ensuite permis de compléter la description du besoin.",
    #             Besoin="L'objectif est de concevoir et fabriquer des prototypes fonctionnels d'un multiplexeur 1=4 permettant l'interconnexion entre le générateur US et une sonde matricielle 1024 voies du système d'échographe ZEUS. ",
    #             Role_TAME ="ICONEUS sollicite Tronico Tame-Care pour l'accompagner dans les différentes étapes d'une conception (HW, Mécanique et firmware) et d'une industrialisation, conformes aux règles de développement et de fabrication d'un sous ensemble devant être intégré dans un dispositif médical."
    #         ), 
    #         hypotheses = [Hypothese(
    #             Type_Hypothese= "Hypothèses ÉlectroniqueS",
    #             List_Hypotheses = "Architecture électronique\n- L'électronique est constituée d'une carte multiplexage principale et de 6 cartes « interposers » identiques\n- Les cartes « interposers » assurent la liaison entre le connecteur DLP408 du générateur ultrason et la carte de multiplexage\n- La carte de multiplexage intègre le connecteur FX11LB-140P-SV(21) assurant l'interface avec la sonde\n\nChaine de multiplexage\n- Réalisation du multiplexage via des commutateurs analogiques pilotés individuellement\n- L'architecture permet la gestion des 1024 voies de la sonde conformément aux configurations reçues\n- Le temps maximum de commutation visé est de 5µs\n\nCommande et supervision\n- Un microcontrôleur assure l'interface de communication SPI avec le générateur ultrason\n- Le microcontrôleur pilote et synchronise les circuits de multiplexage\n- Acquisition des informations de 3 sondes de température\n\nAlimentation\n- L'ensemble des alimentations nécessaires au fonctionnement de la carte est fourni par le générateur ultrason\n- La puissance disponible sur les interfaces d'alimentation est supposée compatible avec les besoins de l'électronique proposée\n- Filtrage des alimentations\n\nSignaux ultrasonores\n- L'architecture est compatible avec des signaux ultrasonores jusqu'à ±100V\n- La conception est dimensionnée pour une fréquence de fonctionnement de 2MHz, avec une évolution future de 15MHz\n- Les voies sont conçues pour respecter une adaptation d'impédance de 50Ω."
    #         )],
    #         lots_list = SyntheseFinanciere(
    #             lots=[Lot(
    #                 phase = "PHASE 1 : PROTOTYPE A",
    #                 nom = "LOT1 = Spécification Technique du Besoin (SR)",
    #                 image_lot = "img/technique.png",
    #                 details_lot ="""Objectif : 
    #     - Consolider les données d'entrée afin d'affiner les spécifications du dispositif selon les besoins utilisateur, du fonctionnement technique et des performances revendiquées par Iconeus
    #     Activités 
    #     - Cahier des charges techniques (TRS)
    #     - Matrice de conformité (CM)
    #     - Initialisation de la liste des composants critiques (LCC)
        
    #     Hors périmètre = Analyse de risque système
        
        
    #     """,
    #                 documents = DocumentsLot(
    #                     Donnees_entree= [Document(img_doc ="img/doc_bleu.png" , text_doc ="Données techniques pour s'interfacer avec la sonde et le générateur US")],
    #                     Livrables = [Document(img_doc ="img/doc_orange.png" , text_doc ="CdC technique (TRS)")]
    #                 ),
    #                 # donnees_tableau={"LOTS":"Lot 0 : Réunion de lancement (KOR)", "DONNÉES D'ENTRÉE":"- Cahier des charges client", "ACTIVITÉS":"- Mise en place des outils de gestion projet - Réunion de lancement (KOM)", "LIVRABLES":"CR KOM", "DÉLAIS":"1 sem", "PRIX DU LOT":"1071", "FACTURATION":"78338"}
    #                 # )]
    #                 debut = 1,
    #                 duree = 2,
    #                 cout = 1000,
    #                 facturation = 5000,
    #                 jalon = [2],
    #                 sous_tache = [SubTache(
    #                     nom = "Tache n°1",
    #                     duree=1,
    #                     debut=1,
    #                     jalon = None,
    #                     couleur="5377B9"
    #                 )])
    #             ]
    #         ),
    #     )

    
    import json

    with open(
        # "input/DEV-26-001521_Devis_PulsHeart_Trolley.json",
        "input/devis_001359_rack_balayage_v2.json",
        "r",
        encoding="utf-8"
    ) as fichier:
        donnees = json.load(fichier)

    test = PresentationData(**donnees)
    # Exemple BETABEAMS


    main(test)

    